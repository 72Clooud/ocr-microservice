import io
import logging
import uuid
from datetime import date, datetime
from typing import Any, Optional

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import Engine, text

from config import settings

logger = logging.getLogger(__name__)

RENAME_COLUMNS = {
    "invoice_id": "ID Faktury",
    "invoice_number": "Numer Faktury",
    "issue_date": "Data Wystawienia",
    "created_at": "Data Utworzenia",
    "currency": "Waluta",
    "net_amount": "Kwota Netto",
    "gross_amount": "Kwota Brutto",
    "status_id": "Status ID",
    "supplier_id": "ID Dostawcy",
    "uploaded_by": "ID Użytkownika",
    "rejection_reason": "Powód Odrzucenia",
}

CENTER_COLUMNS = {
    "ID Faktury",
    "Data Wystawienia",
    "Data Utworzenia",
    "Waluta",
    "Status ID",
    "ID Dostawcy",
    "ID Użytkownika",
}

AMOUNT_COLUMNS = {"Kwota Netto", "Kwota Brutto"}


def _format_worksheet(ws: Any) -> None:
    header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="1F3864")
    cell_font = Font(name="Calibri", size=11, color="1F2937")
    amount_font = Font(name="Calibri", size=11, bold=False, color="111827")

    thin_side = Side(style="thin", color="E2E8F0")
    header_bottom_side = Side(style="medium", color="94A3B8")

    header_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=header_bottom_side,
    )
    cell_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    ws.row_dimensions[1].height = 28

    headers: dict[int, str] = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        header_text = str(cell.value or "")
        headers[col_idx] = header_text
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 21
        row_fill = zebra_fill if row_idx % 2 == 0 else white_fill

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            col_name = headers.get(col_idx, "")

            cell.fill = row_fill
            cell.border = cell_border
            cell.font = cell_font

            if col_name in AMOUNT_COLUMNS:
                cell.font = amount_font
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in CENTER_COLUMNS:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "YYYY-MM-DD"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        col_name = headers.get(col_idx, "")
        max_len = len(col_name)

        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            if isinstance(val, (float, int)) and col_name in AMOUNT_COLUMNS:
                text_val = f"{val:,.2f}"
            elif isinstance(val, datetime):
                text_val = val.strftime("%Y-%m-%d %H:%M")
            elif isinstance(val, date):
                text_val = val.strftime("%Y-%m-%d")
            else:
                text_val = str(val)
            max_len = max(max_len, len(text_val))

        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)


def generate_invoices_report(
    engine: Engine,
    s3_client: Any,
    bucket: Optional[str] = None,
    issue_date_from: Optional[date] = None,
    issue_date_to: Optional[date] = None,
    created_at_from: Optional[date] = None,
    created_at_to: Optional[date] = None,
    status_id: Optional[int] = None,
    uploaded_by: Optional[int] = None,
    supplier_id: Optional[int] = None,
    currency: Optional[str] = None,
    limit: int = 50000,
) -> dict:
    bucket_name = bucket or settings.REPORTS_BUCKET_NAME
    if not bucket_name:
        raise ValueError("MinIO reports bucket name is not configured.")

    if issue_date_from and issue_date_to and issue_date_from > issue_date_to:
        raise ValueError("issue_date_from cannot be later than issue_date_to.")

    if created_at_from and created_at_to and created_at_from > created_at_to:
        raise ValueError("created_at_from cannot be later than created_at_to.")

    clauses = []
    params = {}

    if uploaded_by is not None and uploaded_by > 0:
        clauses.append("uploaded_by = :uploaded_by")
        params["uploaded_by"] = uploaded_by

    if status_id is not None and status_id > 0:
        clauses.append("status_id = :status_id")
        params["status_id"] = status_id

    if supplier_id is not None and supplier_id > 0:
        clauses.append("supplier_id = :supplier_id")
        params["supplier_id"] = supplier_id

    if currency and currency.strip() and currency.strip().lower() not in ("str", "string"):
        clauses.append("currency = :currency")
        params["currency"] = currency.strip().upper()

    if issue_date_from is not None:
        clauses.append("issue_date >= :issue_date_from")
        params["issue_date_from"] = issue_date_from

    if issue_date_to is not None:
        clauses.append("issue_date <= :issue_date_to")
        params["issue_date_to"] = issue_date_to

    if created_at_from is not None:
        clauses.append("created_at >= :created_at_from")
        params["created_at_from"] = created_at_from

    if created_at_to is not None:
        clauses.append("created_at <= :created_at_to")
        params["created_at_to"] = created_at_to

    where_sql = ("WHERE " + " AND ".join(clauses)) if clauses else ""

    params["limit"] = max(1, min(limit, 100000))

    query_str = f"""
        SELECT 
            invoice_id,
            invoice_number,
            issue_date,
            created_at,
            currency,
            net_amount,
            gross_amount,
            status_id,
            supplier_id,
            uploaded_by,
            rejection_reason
        FROM invoices
        {where_sql}
        ORDER BY issue_date DESC NULLS LAST, invoice_id DESC
        LIMIT :limit
    """

    with engine.connect() as conn:
        df = pd.read_sql(text(query_str), conn, params=params)

    export_df = df.rename(columns=RENAME_COLUMNS)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="Faktury")
        _format_worksheet(writer.sheets["Faktury"])
    buffer.seek(0)

    now = datetime.now()
    timestamp_str = now.strftime("%Y%m%d_%H%M%S")
    rand_suffix = uuid.uuid4().hex[:6]
    filename = f"invoices_report_{timestamp_str}_{rand_suffix}.xlsx"
    file_path = f"invoices/{filename}"

    try:
        s3_client.head_bucket(Bucket=bucket_name)
    except Exception:
        try:
            s3_client.create_bucket(Bucket=bucket_name)
        except Exception as e:
            logger.warning(f"[Reports] Could not ensure bucket '{bucket_name}': {e}")

    s3_client.put_object(
        Bucket=bucket_name,
        Key=file_path,
        Body=buffer.getvalue(),
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    logger.info(f"[Reports] Report '{filename}' ({len(df)} rows) uploaded to MinIO at '{bucket_name}/{file_path}'")

    download_url = None
    if settings.MINIO_EXTERNAL_URL:
        download_url = f"{settings.MINIO_EXTERNAL_URL.rstrip('/')}/{bucket_name}/{file_path}"

    return {
        "file_path": file_path,
        "bucket": bucket_name,
        "filename": filename,
        "download_url": download_url,
        "rows_count": len(df),
        "generated_at": now.isoformat(),
    }
